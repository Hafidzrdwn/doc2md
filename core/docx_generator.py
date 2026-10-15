import io
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .themes import get_theme, ThemeConfig
from .docx_helpers import (
    clean_table_cell_text,
    clean_html_formatting,
    add_inline_runs,
    build_docx_table,
    build_docx_code_block,
    set_cell_shading,
    set_cell_margins,
)

def markdown_to_docx_bytes(markdown_text: str, theme: Optional[str] = None) -> bytes:
    cfg = get_theme(theme)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = cfg.docx_body_font
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(*cfg.rgb_text)

    lines = markdown_text.splitlines()
    in_code_block = False
    code_lines = []
    code_lang = ""
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines, in_table
        if table_lines:
            build_docx_table(doc, table_lines, cfg)
        table_lines = []
        in_table = False

    def flush_code_block():
        nonlocal code_lines, in_code_block, code_lang
        if code_lines or in_code_block:
            build_docx_code_block(doc, code_lines, code_lang, cfg)
        code_lines = []
        code_lang = ""
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()

        code_fence_match = re.match(r'^\s*(```|~~~)\s*([a-zA-Z0-9_\-\.\+#]*)\s*$', line)
        if code_fence_match:
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lang = code_fence_match.group(2).strip()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        single_fence_match = re.match(r'^\s*(```|~~~)([a-zA-Z0-9_\-\.\+#]*)\s+(.+?)\s*\1\s*$', line)
        if single_fence_match:
            if in_table:
                flush_table()
            in_code_block = True
            code_lang = single_fence_match.group(2).strip()
            code_lines = [single_fence_match.group(3)]
            flush_code_block()
            i += 1
            continue

        if trimmed.startswith('|') and trimmed.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(trimmed)
            i += 1
            continue
        elif in_table:
            flush_table()

        if not trimmed:
            i += 1
            continue

        h_match = re.match(r'^(#{1,6})\s+(.*)$', trimmed)
        if h_match:
            level = len(h_match.group(1))
            heading_text = h_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True

            sizes = {1: Pt(20), 2: Pt(16), 3: Pt(13), 4: Pt(12), 5: Pt(11), 6: Pt(10.5)}
            heading_rgb = RGBColor(*cfg.rgb_heading)
            colors = {
                1: RGBColor(*cfg.rgb_primary),
                2: heading_rgb,
                3: RGBColor(*cfg.rgb_secondary),
                4: RGBColor(71, 85, 105),
                5: RGBColor(100, 116, 139),
                6: RGBColor(100, 116, 139),
            }
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(4)
            add_inline_runs(p, heading_text)
            for run in p.runs:
                run.bold = True
                run.font.name = cfg.docx_heading_font
                run.font.size = sizes.get(level, Pt(12))
                run.font.color.rgb = colors.get(level, heading_rgb)
            i += 1
            continue

        if trimmed.startswith('>'):
            quote_text = re.sub(r'^>\s?', '', trimmed)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, quote_text)
            for run in p.runs:
                run.italic = True
                run.font.name = cfg.docx_body_font
                run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        if re.match(r'^(?:-{3,}|\*{3,}|_{3,})$', trimmed):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{cfg.table_border_color}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        bullet_match = re.match(r'^(\s*)([-*•]|\+)\s+(.*)$', line)
        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_level = min(indent_spaces // 2, 4)
            item_text = bullet_match.group(3)

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25 * (bullet_level + 1))
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)

            bullet_symbols = ['•', '◦', '▪', '▫', '‣']
            symbol = bullet_symbols[min(bullet_level, len(bullet_symbols) - 1)]

            r_bullet = p.add_run(f"{symbol}  ")
            r_bullet.font.name = 'Arial'
            r_bullet.font.size = Pt(10)
            r_bullet.font.color.rgb = RGBColor(*cfg.rgb_primary)

            add_inline_runs(p, item_text)
            i += 1
            continue

        num_match = re.match(r'^(\s*)(\d+[\.)])\s+(.*)$', line)
        if num_match:
            indent_spaces = len(num_match.group(1))
            num_level = min(indent_spaces // 2, 4)
            num_prefix = num_match.group(2)
            item_text = num_match.group(3)

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35 * (num_level + 1))
            p.paragraph_format.first_line_indent = Inches(-0.35)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            r_num = p.add_run(f"{num_prefix}  ")
            r_num.bold = True
            r_num.font.name = cfg.docx_heading_font
            r_num.font.color.rgb = RGBColor(*cfg.rgb_secondary)

            add_inline_runs(p, item_text)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_runs(p, trimmed)
        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code_block()

    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()
