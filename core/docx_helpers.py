import html
import re
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

def set_cell_shading(cell, color_hex: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top: int = 100, bottom: int = 100, left: int = 140, right: int = 140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def clean_table_cell_text(text: str) -> str:
    if not text:
        return ""
    text = html.unescape(text).replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'([a-zA-Z]{2,})[-‐‑]\s*(?:<br\s*/?>|\n)+\s*([a-zA-Z]{1,})', r'\1\2', text, flags=re.IGNORECASE)
    for _ in range(2):
        text = re.sub(
            r'([a-zA-Z]{2,})\s*(?:<br\s*/?>|\n)+\s*([a-zA-Z]{1,3})(?=\s*[\(\[.,:;?!]|(?:<br|\n|\s|$))',
            r'\1\2', text, flags=re.IGNORECASE
        )
    text = re.sub(r'(?:<br\s*/?>|\n)\s*(?:<br\s*/?>|\n)+', '__PARABREAK__', text, flags=re.IGNORECASE)
    text = re.sub(r'(?:<br\s*/?>|\n)\s*(?=[-*•]|\d+[\.)]\s)', '__LINEBREAK__', text, flags=re.IGNORECASE)
    text = re.sub(r'(?:<br\s*/?>|\n)+', ' ', text, flags=re.IGNORECASE)
    text = text.replace('__PARABREAK__', '\n\n').replace('__LINEBREAK__', '\n')
    return re.sub(r'[ \t]+', ' ', text).strip()

def clean_html_formatting(text: str) -> str:
    if not text:
        return ""
    text = html.unescape(text).replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'([a-zA-Z]{2,})[-‐‑]\s*(?:<br\s*/?>|\n)+\s*([a-zA-Z]{1,})', r'\1\2', text, flags=re.IGNORECASE)
    for _ in range(2):
        text = re.sub(
            r'([a-zA-Z]{2,})\s*(?:<br\s*/?>)+\s*([a-zA-Z]{1,3})(?=\s*[\(\[.,:;?!]|(?:<br|\s|$))',
            r'\1\2', text, flags=re.IGNORECASE
        )
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<(?:b|strong)>(.*?)</(?:b|strong)>', r'**\1**', text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r'<(?:i|em)>(.*?)</(?:i|em)>', r'*\1*', text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r'<(?:s|del|strike)>(.*?)</(?:s|del|strike)>', r'~~\1~~', text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r'<code>(.*?)</code>', r'`\1`', text, flags=re.IGNORECASE | re.DOTALL)
    return re.sub(r'</?[a-zA-Z][^>]*>', '', text)

def add_inline_runs(paragraph, text: str):
    cleaned_text = clean_html_formatting(text)
    sub_lines = cleaned_text.split('\n')
    for line_idx, sub_line in enumerate(sub_lines):
        if line_idx > 0:
            paragraph.add_run().add_break()
        if not sub_line:
            continue
        token_pattern = re.compile(
            r'(\*\*\*.*?\*\*\*|___.*?___|\*\*.*?\*\*|__.*?__|~~.*?~~|(?<!\*)\*[^*]+?\*(?!\*)|(?<!_)_[^_]+?_(?!_)|`[^`]+?`|\[.*?\]\(.*?\))'
        )
        tokens = token_pattern.split(sub_line)
        for token in tokens:
            if not token:
                continue
            if (token.startswith('***') and token.endswith('***')) or (token.startswith('___') and token.endswith('___')):
                run = paragraph.add_run(token[3:-3])
                run.bold, run.italic = True, True
            elif (token.startswith('**') and token.endswith('**')) or (token.startswith('__') and token.endswith('__')):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif (token.startswith('*') and token.endswith('*')) or (token.startswith('_') and token.endswith('_')):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            elif token.startswith('~~') and token.endswith('~~'):
                run = paragraph.add_run(token[2:-2])
                run.font.strike = True
            elif token.startswith('`') and token.endswith('`'):
                run = paragraph.add_run(token[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(190, 24, 93)
            elif token.startswith('[') and '](' in token and token.endswith(')'):
                m = re.match(r'\[(.*?)\]\((.*?)\)', token)
                if m:
                    run = paragraph.add_run(m.group(1))
                    run.underline = True
                    run.font.color.rgb = RGBColor(37, 99, 235)
                else:
                    paragraph.add_run(token)
            else:
                paragraph.add_run(token)

def build_docx_table(doc, table_lines: list, cfg):
    if not table_lines:
        return
    parsed_rows = []
    for tline in table_lines:
        stripped = tline.strip()
        if not stripped:
            continue
        if stripped.startswith('|'):
            stripped = stripped[1:]
        if stripped.endswith('|'):
            stripped = stripped[:-1]
        cols = [clean_table_cell_text(c.strip()) for c in stripped.split('|')]
        if all(re.match(r'^:?-+:?$', c) for c in cols if c):
            continue
        parsed_rows.append(cols)

    if not parsed_rows:
        return

    max_cols = max(len(r) for r in parsed_rows)
    for r in parsed_rows:
        while len(r) < max_cols:
            r.append('')

    total_printable_width = 6.5
    col_weights = []
    for c in range(max_cols):
        lens = [len(r[c]) for r in parsed_rows if c < len(r) and r[c]]
        max_l = max(lens) if lens else 4
        col_weights.append(max(max_l ** 0.75, 3.0))

    sum_weights = sum(col_weights) or 1.0
    col_widths = [(w / sum_weights) * total_printable_width for w in col_weights]
    for idx in range(len(col_widths)):
        if col_widths[idx] < 0.6:
            col_widths[idx] = 0.6
    curr_sum = sum(col_widths) or 1.0
    col_widths = [(w / curr_sum) * total_printable_width for w in col_widths]

    table = doc.add_table(rows=len(parsed_rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = False

    for col_idx, width in enumerate(col_widths):
        table.columns[col_idx].width = Inches(width)

    try:
        table.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    except Exception:
        pass

    for row_idx, row_data in enumerate(parsed_rows):
        row_obj = table.rows[row_idx]
        is_header = (row_idx == 0)
        try:
            row_obj._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        except Exception:
            pass

        for col_idx, raw_cell_value in enumerate(row_data):
            cell = row_obj.cells[col_idx]
            cell.width = Inches(col_widths[col_idx])
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if is_header:
                set_cell_shading(cell, cfg.table_header_bg)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, raw_cell_value)

            if is_header:
                r_th = int(cfg.table_header_text[0:2], 16)
                g_th = int(cfg.table_header_text[2:4], 16)
                b_th = int(cfg.table_header_text[4:6], 16)
                for run in p.runs:
                    run.bold = True
                    run.font.name = cfg.docx_heading_font
                    run.font.color.rgb = RGBColor(r_th, g_th, b_th)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(6)

def build_docx_code_block(doc, code_lines: list, code_lang: str, cfg):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, cfg.code_bg)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    if code_lang:
        p_lang = cell.paragraphs[0]
        p_lang.paragraph_format.space_before = Pt(0)
        p_lang.paragraph_format.space_after = Pt(4)
        r_lang = p_lang.add_run(f"[{code_lang.upper()}]")
        r_lang.font.name = 'Consolas'
        r_lang.font.size = Pt(8.5)
        r_lang.font.color.rgb = RGBColor(100, 116, 139)
        r_lang.bold = True
        p_code = cell.add_paragraph()
    else:
        p_code = cell.paragraphs[0]

    p_code.paragraph_format.space_before = Pt(2)
    p_code.paragraph_format.space_after = Pt(2)
    p_code.paragraph_format.line_spacing = 1.15

    for c_idx, cline in enumerate(code_lines):
        r = p_code.add_run(cline)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        if c_idx < len(code_lines) - 1:
            r.add_break()

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(6)
