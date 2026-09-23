import io
from docx import Document
from core.docx_generator import (
    clean_table_cell_text,
    clean_html_formatting,
    markdown_to_docx_bytes,
)

def test_clean_table_cell_text_rejoins_broken_words():
    raw_header = "Pengalama<br>n<br>(tahun)"
    cleaned = clean_table_cell_text(raw_header)
    assert cleaned == "Pengalaman (tahun)"

def test_clean_table_cell_text_hyphenated_fragment():
    raw_header = "Pengalama-<br>n (tahun)"
    cleaned = clean_table_cell_text(raw_header)
    assert cleaned == "Pengalaman (tahun)"

def test_clean_table_cell_text_preserves_intentional_list():
    raw_cell = "Daftar:<br>- Butir A<br>- Butir B"
    cleaned = clean_table_cell_text(raw_cell)
    assert cleaned == "Daftar:\n- Butir A\n- Butir B"

def test_clean_html_formatting():
    raw = "<b>Tebal</b> dan <i>Miring</i> serta <code>Kode</code>"
    cleaned = clean_html_formatting(raw)
    assert cleaned == "**Tebal** dan *Miring* serta `Kode`"

def test_markdown_to_docx_bytes_valid_document():
    md = """# Judul Dokumen

Ini adalah paragraf pengantar.

| Kandidat | IPK | Tes Teknis | Pengalama<br>n<br>(tahun) | Gaji (jt/bln) |
| --- | --- | --- | --- | --- |
| A | 3.8 | 85 | 3 | 8 |
| B | 3.2 | 75 | 1 | 4.5 |

```python
def hitung():
    return 100
```

1. Kriteria pertama
2. Kriteria kedua

| Nilai | Bobot |
| --- | --- |
| 100 | 50% |

1. Soal nomor satu
2. Soal nomor dua
"""
    docx_bytes = markdown_to_docx_bytes(md)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 0

    doc = Document(io.BytesIO(docx_bytes))
    # Table 0: Main criteria table
    # Table 1: Python code block 1x1 table
    # Table 2: Secondary table
    assert len(doc.tables) == 3

    t0_headers = [c.text.strip() for c in doc.tables[0].rows[0].cells]
    assert t0_headers == ['Kandidat', 'IPK', 'Tes Teknis', 'Pengalaman (tahun)', 'Gaji (jt/bln)']

    # Check proportional column width (Pengalaman (tahun) should have wide allocation ~1.8-2.0 in)
    col3_w = doc.tables[0].columns[3].width.inches
    col1_w = doc.tables[0].columns[1].width.inches
    assert col3_w > col1_w
    assert col3_w > 1.5

    # Check that code block contains badge
    code_cell_text = doc.tables[1].rows[0].cells[0].text
    assert "[PYTHON]" in code_cell_text
    assert "def hitung():" in code_cell_text

    # Check that paragraphs following second table preserve isolated list numbering
    p_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert any(p.startswith("1.  Soal nomor satu") or p.startswith("1. Soal nomor satu") for p in p_texts)
    assert any(p.startswith("2.  Soal nomor dua") or p.startswith("2. Soal nomor dua") for p in p_texts)
